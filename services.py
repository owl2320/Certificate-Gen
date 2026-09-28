"""Pure business logic: no GUI, no web framework. Ported from reader.py / generator.py."""
from collections import defaultdict
from pathlib import Path
from typing import Callable, Optional

import pandas as pd
from bijoy2unicode.converter import Unicode
from docx import Document
from docxcompose.composer import Composer
from docxtpl import DocxTemplate

DATA_EXTENSIONS = {".csv", ".xls", ".xlsx", ".docx"}


# ---------------------------------------------------------------- reading
def read_data(file_path: Path) -> dict[str, list[str]]:
    """Read CSV / Excel / DOCX into {column: [values...]} (all values as str)."""
    suffix = file_path.suffix.lower()

    if suffix in (".csv", ".xls", ".xlsx"):
        # dtype=str keeps leading zeros / IDs intact; keep_default_na avoids "nan"
        if suffix == ".csv":
            df = pd.read_csv(file_path, dtype=str, keep_default_na=False)
        else:
            df = pd.read_excel(file_path, dtype=str, keep_default_na=False)
        df.columns = [str(c).strip() for c in df.columns]
        return {col: df[col].tolist() for col in df.columns}

    if suffix == ".docx":
        doc = Document(file_path)
        headers: Optional[list[str]] = None
        data: dict[str, list[str]] = defaultdict(list)
        for table in doc.tables:
            for row in table.rows:
                texts = [cell.text.strip() for cell in row.cells]
                if headers is None:
                    headers = texts
                    continue
                if texts == headers:  # repeated header row across pages/tables
                    continue
                for i, h in enumerate(headers):
                    data[h].append(texts[i] if i < len(texts) else "")
        return dict(data)

    raise ValueError(f"Unsupported file type: {suffix}")


def convert_unicode(data: dict[str, list]) -> dict[str, list[str]]:
    """Bijoy -> Unicode for both headers and values."""
    conv = Unicode()
    return {
        conv.convertBijoyToUnicode(str(k)): [conv.convertBijoyToUnicode(str(v)) for v in vals]
        for k, vals in data.items()
    }


def get_data(file_path: Path, is_bijoy: bool) -> dict[str, list[str]]:
    raw = read_data(file_path)
    return convert_unicode(raw) if is_bijoy else raw


# ------------------------------------------------------------- generating
def generate_certs(
    template_path: Path,
    data: dict[str, list],
    output_dir: Path,
    on_progress: Optional[Callable[[int, int], None]] = None,
) -> list[Path]:
    """Render one .docx per row. Raises on failure (caller reports it)."""
    output_dir.mkdir(parents=True, exist_ok=True)
    total = min(len(v) for v in data.values()) if data else 0
    files: list[Path] = []

    for i in range(total):
        try:
            doc = DocxTemplate(template_path)  # must be fresh per render
            doc.render({k: v[i] for k, v in data.items()})
            out = output_dir / f"{i + 1}.docx"
            doc.save(out)
        except Exception as e:
            raise RuntimeError(f"Error on row {i + 1}: {e}") from e
        files.append(out)
        if on_progress:
            on_progress(i + 1, total)
    return files


def merge_docs(files: list[Path], output_file: Path) -> Path:
    """Merge docx files in the order given (already row-ordered, no regex sorting needed)."""
    if not files:
        raise ValueError("No .docx files to merge")
    composer = Composer(Document(files[0]))
    for f in files[1:]:
        composer.append(Document(f))
    composer.save(output_file)
    return output_file